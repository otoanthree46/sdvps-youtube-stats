"""old_tw.xlsx(1行目=日付、B=選手名、C=XのURL)のXフォロワー履歴を取り込む。
- 名前で channels.csv と照合(改名: Ryu→DarknessMare)。x_handle が空なら URL から補完
- 既存の (date,key) は上書きしない(食い違いは件数を表示)
usage: python -I scripts/import_x_history.py <xlsx path>
"""
import csv, datetime as dt, sys
import openpyxl

ALIAS = {"Ryu": "DarknessMare"}
CH, XF = "docs/data/channels.csv", "docs/data/x_followers.csv"
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
dcols = [(i, v.date().isoformat()) for i, v in enumerate(rows[0]) if isinstance(v, dt.datetime)]

chans = list(csv.DictReader(open(CH, encoding="utf-8-sig")))
fields = list(chans[0].keys())
existing = {(r["date"], r["key"]): int(r["followers"]) for r in csv.DictReader(open(XF, encoding="utf-8-sig"))}

out, miss, conflict = [], [], 0
for r in rows[1:]:
    name = ALIAS.get(r[1], r[1])
    c = next((c for c in chans if c["name"] == name), None)
    if not c:
        miss.append(name); continue
    h = (r[2] or "").rstrip("/").split("/")[-1].split("?")[0]
    if h and not c["x_handle"]: c["x_handle"] = h
    for i, d in dcols:
        v = r[i] if i < len(r) else None
        if not isinstance(v, (int, float)): continue
        k = (d, c["key"])
        if k in existing:
            conflict += existing[k] != int(v)
        else:
            existing[k] = int(v); out.append([d, c["key"], int(v), "excel"])

with open(CH, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(chans)
old = list(csv.reader(open(XF, encoding="utf-8-sig")))[1:]
allr = sorted(old + [[str(x) for x in o] for o in out], key=lambda x: (x[0], x[1]))
with open(XF, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["date", "key", "followers", "source"]); w.writerows(allr)
print("added:", len(out), "| not found:", miss, "| same date but different value (kept existing):", conflict)